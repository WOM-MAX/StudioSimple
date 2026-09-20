import openpyxl
import json
import os
import re

excel_path = r"d:\StudioSimple - Antigravity\Base-Excel\planes_consolidados_master_enriquecido.xlsx"
output_dir = r"d:\StudioSimple - Antigravity\PROTOTIPO\prototipo\public\data"
os.makedirs(output_dir, exist_ok=True)
output_json = os.path.join(output_dir, "curriculum_catalog.json")

print(f"Loading Excel from {excel_path}...")
wb = openpyxl.load_workbook(excel_path, read_only=True)
sheet = wb.active

header = [str(cell or '').strip() for cell in next(sheet.iter_rows(max_row=1, values_only=True))]
col_map = {name: idx for idx, name in enumerate(header)}

# Official MINEDUC Exámenes Libres 7° Básico Temario Map (OAs in order of the official document)
OFFICIAL_TEMARIO_EELL_7B = {
    "Matemática": [1, 3, 4, 6, 8, 11, 14, 16, 18],
    "Lengua y Literatura": [3, 4, 9, 15],
    "Ciencias Naturales": [1, 2, 5, 7, 9, 13, 14],
    "Historia, Geografía y Ciencias Sociales": [2, 3, 6, 7, 9, 12, 13, 16, 18, 19, 20, 21, 22],
    "Inglés": [9, 10, 13, 16]
}

def determine_lessons(desc, indicadores, asignatura, oa_num):
    text = (desc + " " + " ".join(indicadores)).lower()
    
    dual_patterns = [
        r"multiplicaci[oó]n y divisi[oó]n",
        r"adici[oó]n y sustracci[oó]n",
        r"per[ií]metro y [aá]rea",
        r"directa e inversa",
        r"virus.*bacteria",
        r"biol[oó]gicos.*sociales",
        r"causas y consecuencias",
    ]
    
    for pat in dual_patterns:
        if re.search(pat, text):
            return 6, "Presenta dualidad de procesos complejos o habilidades independientes que requieren modelado por separado."
            
    if len(indicadores) >= 6:
        return 6, f"Contiene {len(indicadores)} indicadores de evaluación oficiales que exceden la carga cognitiva de 5 clases de 30 min."
        
    return 5, "Secuencia estándar de 5 clases: activación, modelo visual, procedimiento 1, procedimiento 2 y síntesis con problemas."

catalog = []
seen = set()

target_subjects = {
    "Matemática": "Matemática",
    "Ciencias Naturales": "Ciencias Naturales",
    "Historia, Geografía y Ciencias Sociales": "Historia, Geografía y Ciencias Sociales",
    "Lengua y Literatura": "Lengua y Literatura",
    "Lenguaje": "Lengua y Literatura",
    "Idioma Extranjero Inglés": "Inglés",
    "Inglés": "Inglés"
}

for row in sheet.iter_rows(min_row=2, values_only=True):
    curso = str(row[col_map.get('Curso', 0)] or '').strip()
    asig_raw = str(row[col_map.get('Asignatura', 1)] or '').strip()
    eje = str(row[col_map.get('Eje Curricular', 2)] or '').strip()
    oa_str = str(row[col_map.get('N° de OA', 3)] or '').strip()
    desc = str(row[col_map.get('Descripción del OA', 4)] or row[4] or '').strip()
    indicadores_raw = str(row[col_map.get('Indicadores de Evaluación', 6)] or row[6] or '').strip()
    conceptos_raw = str(row[col_map.get('Conceptos_Clave', 13)] or row[13] or '').strip()
    
    if not oa_str or not desc or curso != "7° Básico":
        continue
        
    asig_norm = None
    for k, v in target_subjects.items():
        if k.lower() in asig_raw.lower():
            asig_norm = v
            break
            
    if not asig_norm:
        continue
        
    unique_key = (curso, asig_norm, oa_str)
    if unique_key in seen:
        continue
    seen.add(unique_key)
    
    # Process indicators
    indicators = [
        re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()
        for line in indicadores_raw.splitlines()
        if len(line.strip()) > 5
    ]
    
    concepts = [
        c.strip() for c in re.split(r'[,;\n•]', conceptos_raw)
        if len(c.strip()) > 2
    ]
    
    oa_match = re.search(r'\d+', oa_str)
    oa_order = int(oa_match.group()) if oa_match else 99
    
    lessons_count, rationale = determine_lessons(desc, indicators, asig_norm, oa_order)
    
    # Check official Temario EELL
    temario_list = OFFICIAL_TEMARIO_EELL_7B.get(asig_norm, [])
    in_temario = oa_order in temario_list
    temario_pos = temario_list.index(oa_order) + 1 if in_temario else None
    
    # The first 2 OAs in the official EELL temario are PRIORITY DEMO for funding!
    is_priority = in_temario and (temario_pos in [1, 2])
    
    catalog.append({
        "id": f"110-7-{asig_norm[:3].upper()}-OA{oa_order:02d}",
        "curso": curso,
        "asignatura": asig_norm,
        "eje": eje,
        "oa": f"OA {oa_order}",
        "oaNumero": oa_order,
        "inTemarioEELL": in_temario,
        "temarioPosicion": temario_pos,
        "isPriorityDemo": is_priority,
        "descripcion": desc,
        "indicadores": indicators,
        "conceptosClave": concepts,
        "leccionesSugeridas": lessons_count,
        "justificacionLecciones": rationale
    })

# Sort by Subject, Temario presence, then OA Number
catalog.sort(key=lambda x: (
    x["asignatura"],
    0 if x["inTemarioEELL"] else 1,
    x["temarioPosicion"] if x["temarioPosicion"] else 999,
    x["oaNumero"]
))

print(f"Total OAs extracted for 7° Básico: {len(catalog)}")
with open(output_json, "w", encoding="utf-8") as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)

print(f"Catalog successfully updated in {output_json}")
