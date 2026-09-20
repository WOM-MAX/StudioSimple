import openpyxl
import json
import os
import re

excel_candidates = [
    r"E:\Proyectos\app-edutech\BASES DE DATOS\planes_consolidados_master_enriquecido.xlsx",
    r"d:\StudioSimple - Antigravity\Base-Excel\planes_consolidados_master_enriquecido.xlsx"
]

out_paths = [
    r"d:\StudioSimple - Antigravity\Web Studio Simple\public\data\curriculum_catalog.json",
    r"d:\StudioSimple - Antigravity\PROTOTIPO\prototipo\public\data\curriculum_catalog.json"
]

target_subjects = {
    "Matemática": "Matemática",
    "Ciencias Naturales": "Ciencias Naturales",
    "Historia, Geografía y Ciencias Sociales": "Historia, Geografía y Ciencias Sociales",
    "Lengua y Literatura": "Lengua y Literatura",
    "Lenguaje": "Lengua y Literatura",
    "Idioma Extranjero Inglés": "Inglés",
    "Inglés": "Inglés"
}

# Temarios oficiales extraídos directamente de los documentos MINEDUC en TEMARIOS EELL/ (3° a 8° Básico)
OFFICIAL_TEMARIO_EELL = {
    "3° Básico": {
        "Lengua y Literatura": [4, 5, 6, 17],
        "Matemática": [2, 3, 5, 6, 8, 9, 10, 11, 12, 15, 21, 22, 25],
        "Ciencias Naturales": [2, 4, 5, 6, 10, 11, 12],
        "Historia, Geografía y Ciencias Sociales": [3, 4, 7, 8, 14]
    },
    "4° Básico": {
        "Lengua y Literatura": [4, 5, 6, 17],
        "Matemática": [1, 3, 5, 6, 7, 8, 9, 13, 14, 17, 18, 22, 23, 27],
        "Ciencias Naturales": [1, 2, 3, 6, 7, 10, 15, 16],
        "Historia, Geografía y Ciencias Sociales": [4, 8, 9, 10, 11, 12]
    },
    "5° Básico": {
        "Lengua y Literatura": [4, 5, 6, 7, 14, 15],
        "Matemática": [1, 3, 4, 6, 7, 10, 11, 13, 14, 15, 17, 18, 19, 22, 23, 24, 26],
        "Ciencias Naturales": [1, 5, 6, 7, 10, 12, 14],
        "Historia, Geografía y Ciencias Sociales": [2, 3, 6, 7, 9, 10, 13, 14],
        "Inglés": [5, 6, 8, 14, 15]
    },
    "6° Básico": {
        "Lengua y Literatura": [4, 5, 6, 7, 14, 15],
        "Matemática": [2, 3, 4, 5, 8, 11, 16, 18, 19, 23, 24],
        "Ciencias Naturales": [1, 2, 5, 6, 8, 11, 13, 16, 18],
        "Historia, Geografía y Ciencias Sociales": [1, 2, 3, 5, 7, 8, 9, 11, 12, 15, 16, 17, 18],
        "Inglés": [5, 6, 8, 14, 15]
    },
    "7° Básico": {
        "Lengua y Literatura": [3, 4, 7, 8, 9, 15],
        "Matemática": [1, 3, 4, 6, 8, 11, 14, 16, 18],
        "Ciencias Naturales": [1, 2, 5, 7, 9, 13, 14],
        "Historia, Geografía y Ciencias Sociales": [2, 3, 6, 7, 9, 12, 13, 16, 18, 19, 20, 21, 22],
        "Inglés": [9, 10, 13, 16]
    },
    "8° Básico": {
        "Lengua y Literatura": [3, 4, 8, 9, 10, 16],
        "Matemática": [1, 2, 3, 4, 8, 10, 12, 13, 15, 16],
        "Ciencias Naturales": [2, 4, 5, 7, 10, 11, 12, 14],
        "Historia, Geografía y Ciencias Sociales": [2, 3, 4, 7, 11, 12, 14, 16, 18, 20, 21, 22],
        "Inglés": [9, 10, 13, 16]
    }
}

def determine_lessons(desc, indicadores):
    text = (desc + " " + " ".join(indicadores)).lower()
    dual_patterns = [
        r"multiplicaci[oó]n y divisi[oó]n",
        r"adici[oó]n y sustracci[oó]n",
        r"per[ií]metro y [aá]rea",
        r"directa e inversa",
        r"virus.*bacteria",
        r"biol[oó]gicos.*sociales",
        r"causas y consecuencias",
    ]
    for pat in dual_patterns:
        if re.search(pat, text):
            return 6, "Presenta dualidad de procesos complejos o habilidades independientes que requieren modelado por separado."
            
    if len(indicadores) >= 6:
        return 6, f"Contiene {len(indicadores)} indicadores de evaluación oficiales que exceden la carga cognitiva de 5 clases de 30 min."
        
    return 5, "Secuencia estándar de 5 clases: activación, modelo visual, procedimiento 1, procedimiento 2 y síntesis con problemas."

excel_path = None
for candidate in excel_candidates:
    if os.path.exists(candidate):
        excel_path = candidate
        break

catalog = []

if excel_path:
    print(f"Cargando Excel desde {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    sheet = wb.active

    header = [str(cell or '').strip() for cell in next(sheet.iter_rows(max_row=1, values_only=True))]
    col_map = {name: idx for idx, name in enumerate(header)}

    seen = set()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        curso_raw = str(row[col_map.get('Curso', 0)] or '').strip()
        m = re.search(r'([3-8])', curso_raw)
        if not m or 'Medio' in curso_raw:
            continue
            
        grade_num = int(m.group(1))
        curso_norm = f"{grade_num}° Básico"
        
        asig_raw = str(row[col_map.get('Asignatura', 1)] or '').strip()
        eje = str(row[col_map.get('Eje Curricular', 2)] or '').strip()
        oa_str = str(row[col_map.get('N° de OA', 3)] or '').strip()
        desc = str(row[col_map.get('Descripción del OA', 4)] or row[4] or '').strip()
        indicadores_raw = str(row[col_map.get('Indicadores de Evaluación', 6)] or row[6] or '').strip()
        conceptos_raw = str(row[col_map.get('Conceptos_Clave', 13)] or row[13] or '').strip()
        
        if not oa_str or not desc:
            continue
            
        asig_norm = None
        for k, v in target_subjects.items():
            if k.lower() in asig_raw.lower():
                asig_norm = v
                break
                
        if not asig_norm:
            continue
            
        oa_match = re.search(r'\d+', oa_str)
        if not oa_match:
            continue
        oa_order = int(oa_match.group())

        # FILTRO ESTRICTO: Solo OAs evaluables en Exámenes Libres MINEDUC
        temario_dict = OFFICIAL_TEMARIO_EELL.get(curso_norm, {})
        temario_list = temario_dict.get(asig_norm, [])
        if not temario_list or oa_order not in temario_list:
            continue
            
        unique_key = (curso_norm, asig_norm, oa_order)
        if unique_key in seen:
            continue
        seen.add(unique_key)
        
        indicators = [
            re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()
            for line in indicadores_raw.splitlines()
            if len(line.strip()) > 5
        ]
        
        concepts = [
            c.strip() for c in re.split(r'[,;\n•]', conceptos_raw)
            if len(c.strip()) > 2
        ]
        
        lessons_count, rationale = determine_lessons(desc, indicators)
        
        temario_pos = temario_list.index(oa_order) + 1
        is_priority = temario_pos in [1, 2]
        
        asig_code = asig_norm[:3].upper()
        if asig_norm == "Historia, Geografía y Ciencias Sociales":
            asig_code = "HIS"
        elif asig_norm == "Lengua y Literatura":
            asig_code = "LEN"
        elif asig_norm == "Ciencias Naturales":
            asig_code = "CIE"
        elif asig_norm == "Matemática":
            asig_code = "MAT"
        elif asig_norm == "Inglés":
            asig_code = "ING"
            
        oa_id = f"110-{grade_num}-{asig_code}-OA{oa_order:02d}"
        
        catalog.append({
            "id": oa_id,
            "curso": curso_norm,
            "asignatura": asig_norm,
            "eje": eje,
            "oa": f"OA {oa_order}",
            "oaNumero": oa_order,
            "inTemarioEELL": True,
            "temarioPosicion": temario_pos,
            "isPriorityDemo": is_priority,
            "descripcion": desc,
            "indicadores": indicators,
            "conceptosClave": concepts,
            "leccionesSugeridas": lessons_count,
            "justificacionLecciones": rationale
        })
else:
    print("Archivo Excel no encontrado, filtrando desde curriculum_catalog.json existente...")
    source_json = out_paths[0]
    with open(source_json, "r", encoding="utf-8") as f:
        full_cat = json.load(f)
    
    for item in full_cat:
        curso_norm = item["curso"]
        asig_norm = item["asignatura"]
        oa_order = item["oaNumero"]
        temario_dict = OFFICIAL_TEMARIO_EELL.get(curso_norm, {})
        temario_list = temario_dict.get(asig_norm, [])
        if temario_list and oa_order in temario_list:
            pos = temario_list.index(oa_order) + 1
            item_copy = dict(item)
            item_copy["inTemarioEELL"] = True
            item_copy["temarioPosicion"] = pos
            item_copy["isPriorityDemo"] = pos in [1, 2]
            catalog.append(item_copy)

catalog.sort(key=lambda x: (
    x["curso"],
    x["asignatura"],
    x["temarioPosicion"] if x["temarioPosicion"] else 999,
    x["oaNumero"]
))

print(f"Total OAs oficiales procesados para Exámenes Libres (3° a 8° Básico): {len(catalog)}")

for p in out_paths:
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(catalog, f, ensure_ascii=False, indent=2)
    print(f"Archivo guardado exitosamente en: {p}")
