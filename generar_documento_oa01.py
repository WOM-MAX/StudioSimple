# -*- coding: utf-8 -*-
"""
Script Completo para Generación del Documento Word Oficial de Lecciones OA 01
Asignatura: Matemática - 7° Básico
Objetivo de Aprendizaje: OA 01 (Números Enteros Z)
Total: 5 Lecciones Completas + 10 Guiones de Video Detallados
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_oa01_word_doc(output_path):
    doc = docx.Document()
    
    # Configuración de márgenes (Estándar 2.0 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def add_callout_box(title, text, bg_color="F8FAFC", border_color="0D9488", title_color="0F766E"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        run_t = p.add_run(f"[{title}]\n")
        run_t.bold = True
        run_t.font.name = "Calibri"
        run_t.font.size = Pt(10)
        run_t.font.color.rgb = RGBColor.from_string(title_color)
        
        run_b = p.add_run(text)
        run_b.font.name = "Calibri"
        run_b.font.size = Pt(9.5)
        run_b.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(3)

    def add_video_script_table(script_data):
        table = doc.add_table(rows=len(script_data) + 1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        headers = ["Escena", "Tiempo", "Visual y Animación", "Locución (Voz en Off)", "Texto en Pantalla"]
        col_widths = [Inches(0.7), Inches(0.7), Inches(1.9), Inches(2.3), Inches(1.2)]
        
        hdr_row = table.rows[0]
        for i, h_text in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, "1E293B")
            set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h_text)
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        for row_idx, item in enumerate(script_data):
            row = table.rows[row_idx + 1]
            bg = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
            for col_idx, key in enumerate(["escena", "tiempo", "visual", "locucion", "texto_pantalla"]):
                cell = row.cells[col_idx]
                cell.width = col_widths[col_idx]
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
                p = cell.paragraphs[0]
                if col_idx in [0, 1]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(item[key])
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(4)

    # ==========================================
    # PORTADA DEL DOCUMENTO
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(12)
    
    r_main = title_p.add_run("STUDIOSIMPLE - ECOSISTEMA EDUCATIVO\n")
    r_main.bold = True
    r_main.font.name = "Calibri"
    r_main.font.size = Pt(22)
    r_main.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r_sub = title_p.add_run("PROGRAMA DE LECCIONES COMPLETAS Y GUIONES AUDIOVISUALES\nMATEMÁTICA 7° BÁSICO - OBJETIVO DE APRENDIZAJE 01")
    r_sub.bold = True
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    r_meta = p_meta.add_run("Estructura Pedagógica de Sincronización Dual (Mentor / Estudiante)\n5 Lecciones de 30 Minutos | 10 Guiones de Video Detallados (Gancho y Formalización)\nBasado en el Marco Curricular Nacional y Textos Escolares Digitales")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    doc.add_page_break()

    # ==========================================
    # ÍNDICE GENERAL DEL OA 01
    # ==========================================
    h_idx = doc.add_heading("Índice Curricular del OA 01: Números Enteros (Z)", level=1)
    h_idx.paragraph_format.space_after = Pt(12)
    
    idx_text = (
        "El presente documento consolida la planificación instruccional completa del Objetivo de Aprendizaje 01 de 7° Básico: "
        "'Mostrar que comprenden los números enteros: representando en la recta numérica, comparando y ordenando, dando significado a los signos + y -, y resolviendo problemas de la vida diaria'.\n\n"
        "Cada lección está estructurada en el modelo de 8 fases de StudioSimple, garantizando un andamiaje socrático no punitivo para el adulto mediador y dos piezas audiovisuales clave por clase (Gancho Multimedia y Formalización Conceptual)."
    )
    p_idx = doc.add_paragraph(idx_text)
    p_idx.paragraph_format.space_after = Pt(14)
    
    # Tabla resumen de las 5 clases
    resumen_table = doc.add_table(rows=6, cols=4)
    resumen_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_headers = ["Clase", "Título de la Lección", "Foco Didáctico Principal", "Videos Producidos"]
    r_widths = [Inches(0.8), Inches(2.2), Inches(2.6), Inches(1.2)]
    
    for i, h in enumerate(r_headers):
        c = resumen_table.rows[0].cells[i]
        c.width = r_widths[i]
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 80, 80, 80, 80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    resumen_data = [
        ("Clase 1", "Posiciones respecto de un punto de referencia", "El número 0 como origen, temperaturas bajo cero, niveles de profundidad marina y niveles de edificio.", "Gancho (Submarino)\nFormalización"),
        ("Clase 2", "La recta numérica y orden en Z", "Representación horizontal de Z, convenio derecha (+) e izquierda (-), criterio de orden 'más a la derecha'.", "Gancho (Cañón Polar)\nFormalización"),
        ("Clase 3", "Valor absoluto y números opuestos", "Comprensión geométrica del valor absoluto como distancia pura al cero y simetría de opuestos aditivos.", "Gancho (Drones)\nFormalización"),
        ("Clase 4", "Adición y sustracción en Z", "Desplazamientos direccionales, adición de igual/distinto signo y la resta como adición del opuesto aditivo.", "Gancho (Robot Marte)\nFormalización"),
        ("Clase 5", "Resolución de problemas cotidianos y síntesis", "Modelamiento en 4 pasos (Pólya) para saldos bancarios, variaciones térmicas y evaluación sumativa.", "Gancho (Base Antártica)\nFormalización")
    ]
    
    for row_idx, r_item in enumerate(resumen_data):
        row = resumen_table.rows[row_idx + 1]
        bg = "FFFFFF" if row_idx % 2 == 0 else "F1F5F9"
        for col_idx, val in enumerate(r_item):
            c = row.cells[col_idx]
            c.width = r_widths[col_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 80, 80)
            p = c.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            
    doc.add_page_break()

    # =========================================================================
    # DETALLE DE CADA UNA DE LAS 5 CLASES
    # =========================================================================
    
    lessons_data = [
        # -------------------------------------------------------------
        # CLASE 1
        # -------------------------------------------------------------
        {
            "num": 1,
            "title": "Posiciones respecto de un punto de referencia",
            "oa": "OA 01 - Clase 1 de 5",
            "duracion": "30 Minutos",
            "objetivo_adulto": "Guiar al estudiante para que comprenda que los números enteros representan posiciones respecto de un punto de referencia y que el cero define ese origen.",
            "ruta_sesion": "Termómetro y ascensor → Video Desafío del submarino → Conversación guiada → Explicación formal → Práctica contextual → Miniquiz → Cierre metacognitivo.",
            "mentor_tip": "Lee únicamente los recuadros 'DILE' en voz alta. Si tu hijo duda, utiliza el apoyo socrático sin dar la respuesta de inmediato.",
            "clima_emocional": "Crea un clima seguro: 'Aquí no buscamos rapidez; equivocarse es la forma en que el cerebro descubre nuevas pistas'.",
            "situacion_intro": "Observa este termómetro. El cero es nuestro punto de referencia. Si la temperatura está tres grados sobre cero, escribimos +3 °C. ¿Qué número usaríamos para representar una temperatura de tres grados bajo cero?",
            "resp_esperada_sit": "−3 °C (o menos 3 grados)",
            "pista_socratica_sit": "Si duda o dice solo 3: 'Si a 3 sobre cero le ponemos +3, ¿cómo diferenciamos en el papel que está hacia el otro lado del cero?'",
            "referencia_texto": "En este termómetro usamos el 0 como punto de referencia. Las temperaturas sobre 0 son positivas; las que están bajo 0, son negativas. ¿Qué representa exactamente −3 °C?",
            
            # VIDEO 1: GANCHO
            "v1_titulo": "Video 1 (Gancho H.O.O.K.): La Expedición del Submarino Abisal",
            "v1_dile_antes": "Antes de ver el video, observa con atención dónde inicia el submarino y fíjate bien en cuánto baja y cuánto sube.",
            "v1_haz": "Reproducir video del submarino en la pantalla del estudiante.",
            "v1_script": [
                {"escena": "1", "tiempo": "0:00 - 0:08", "visual": "Plano general de la superficie del mar en calma con cielo despejado. Línea de agua brillante.", "locucion": "A más de veinte metros bajo el mar, una pequeña expedición científica se acerca a una zona inexplorada.", "texto_pantalla": "Superficie del mar (Nivel 0 m)"},
                {"escena": "2", "tiempo": "0:08 - 0:16", "visual": "La cámara baja bajo el agua. El submarino desciende lentamente hasta estabilizarse.", "locucion": "El submarino desciende hasta situarse a veinte metros bajo la superficie.", "texto_pantalla": "Profundidad: −20 m"},
                {"escena": "3", "tiempo": "0:16 - 0:25", "visual": "Pantalla de radar en la cabina parpadeando. El submarino enciende focos potentes y desciende más.", "locucion": "De pronto, el radar detecta una señal misteriosa: para llegar a ella deben descender quince metros más.", "texto_pantalla": "Descenso: Bajar 15 m"},
                {"escena": "4", "tiempo": "0:25 - 0:33", "visual": "Aparece una formación rocosa y arrecife afilado en el fondo marino frente a la proa.", "locucion": "Justo cuando están por alcanzar el objetivo, aparece una barrera rocosa que bloquea el paso.", "texto_pantalla": "¡Obstáculo en el fondo!"},
                {"escena": "5", "tiempo": "0:33 - 0:42", "visual": "Los propulsores del submarino se orientan hacia arriba y la nave asciende con suavidad.", "locucion": "Para evitarla, deben subir ocho metros. Cada movimiento cambia la posición del submarino.", "texto_pantalla": "Ascenso: Subir 8 m"},
                {"escena": "6", "tiempo": "0:42 - 0:50", "visual": "El submarino flota suspendido con instrumentos titilando en la penumbra marina.", "locucion": "Ahora surge el verdadero desafío: tras bajar y luego subir, ¿a qué profundidad quedaron finalmente?", "texto_pantalla": "¿Posición final = ?"},
                {"escena": "7", "tiempo": "0:50 - 0:58", "visual": "Icono gráfico de StudioSimple con una recta vertical y el número 0 parpadeando en la superficie.", "locucion": "Hay una forma matemática de calcularlo sin perderse. ¡Descubrámoslo en la lección!", "texto_pantalla": "StudioSimple Matemática"}
            ],
            "v1_dile_despues": "Excelente observación. Ahora conversaremos sobre lo que vimos en el recorrido del submarino.",
            
            # CONVERSACIÓN SOCRÁTICA
            "conv_items": [
                ("¿Dónde está el punto cero en la historia del submarino?", "En la superficie del mar.", "Si dice 'en el agua', pregunta: '¿Desde qué límite exacto empezamos a medir la profundidad?'"),
                ("¿Qué significa que el submarino esté a −20 metros?", "Es una posición fija: 20 metros bajo la superficie.", "Pregunta: '¿Esa frase dice dónde se encuentra quieto o cómo se traslada?'"),
                ("“Bajar 15 metros”, ¿indica una posición o un movimiento?", "Un movimiento (acción de desplazarse hacia abajo).", "Resalta el verbo: '¿Bajar es estar en un lugar o es trasladarse de un sitio a otro?'"),
                ("“Subir 8 metros”, ¿indica una posición o un movimiento?", "Un movimiento (desplazamiento hacia arriba).", "Pídele que compare: '¿Subir es el punto de llegada o la acción de ascender?'")
            ],
            
            # VIDEO 2: FORMALIZACIÓN
            "v2_titulo": "Video 2 (Formalización): Posición vs Movimiento y el Origen Cero",
            "v2_dile_antes": "Ya descubrimos que dónde está algo es distinto a cómo se mueve. Ahora veremos cómo la matemática representa formalmente estas ideas.",
            "v2_haz": "Reproducir explicación visual de formalización.",
            "v2_script": [
                {"escena": "1", "tiempo": "0:00 - 0:10", "visual": "Gráfico animado con una línea horizontal y vertical que se cruzan en un círculo brillante rotulado '0'.", "locucion": "En matemáticas, para saber dónde está cualquier objeto necesitamos siempre un punto de partida: el punto de referencia cero.", "texto_pantalla": "Punto de Referencia (0)"},
                {"escena": "2", "tiempo": "0:10 - 0:22", "visual": "Aparece un edificio con niveles sobre el suelo (+1, +2, +3) y subterráneos (−1, −2, −3). El nivel de calle es 0.", "locucion": "Todo lo que se ubica por encima o a favor del origen lleva signo positivo (+). Todo lo que se ubica por debajo o en sentido opuesto lleva signo negativo (−).", "texto_pantalla": "Posiciones: +3 (Piso 3) | −2 (Subterráneo 2)"},
                {"escena": "3", "tiempo": "0:22 - 0:35", "visual": "Animación comparativa dividida en dos columnas: Izquierda 'Posición' (punto fijo con chincheta), Derecha 'Movimiento' (flecha dinámica).", "locucion": "¡Atención a esta diferencia clave! Una Posición nos dice dónde está algo de forma fija. Un Movimiento describe una acción o traslado.", "texto_pantalla": "Posición = ¿Dónde está? (ej. −5 m)\nMovimiento = ¿Cómo cambia? (ej. Bajar 3 m)"},
                {"escena": "4", "tiempo": "0:35 - 0:48", "visual": "Un termómetro interactivo marca −4 °C y sube 6 °C, llegando a +2 °C con números destacados.", "locucion": "Así, estar a menos cuatro grados es una posición. Si la temperatura sube seis grados, ese movimiento nos lleva a una nueva posición: más dos grados.", "texto_pantalla": "−4 °C + Subir 6 °C = +2 °C"},
                {"escena": "5", "tiempo": "0:48 - 0:55", "visual": "Resumen con los símbolos clave de Z: {..., −3, −2, −1, 0, +1, +2, +3, ...}.", "locucion": "Los números enteros nos permiten describir el mundo real con total precisión.", "texto_pantalla": "Conjunto Z: Enteros Positivos, Cero y Negativos"}
            ],
            "idea_clave": "Un número entero representa una posición respecto de un punto de referencia (el cero). Una posición indica dónde está algo; un movimiento indica cómo cambia de lugar.",
            
            # PRÁCTICA CONJUNTA
            "practica": [
                ("Situación 1: Temperatura en la cordillera", "El termómetro marca 4 grados bajo cero. ¿Qué número entero representa esta temperatura?", "−4 °C", "Pídele ubicar el cero como referencia. Como está debajo, se usa el signo negativo."),
                ("Situación 2: Ascensor de centro comercial", "Un ascensor baja 5 pisos. ¿Esta frase describe una posición o un movimiento?", "Un movimiento.", "Pregunta si la frase dice en qué piso se detuvo o qué acción realizó."),
                ("Situación 3: Saldo bancario", "Una cuenta corriente tiene una deuda de $12.000. ¿Cómo se expresa con un número entero?", "−$12.000 (o −12000)", "El saldo a favor es positivo (+); tener deuda está por debajo del cero financiero (−).")
            ],
            
            # MINIQUIZ
            "quiz": [
                ("1. Un buzo está sumergido a 14 metros bajo el nivel del mar. ¿Qué número entero representa su posición?", ["A) +14 m", "B) −14 m", "C) 0 m", "D) 14 m sin signo"], "B", "El nivel del mar es 0. Al estar por debajo, se representa con signo negativo: −14 m."),
                ("2. ¿Cuál de las siguientes expresiones representa un MOVIMIENTO y no una posición fija?", ["A) El avión vuela a +3.000 metros.", "B) El congelador está a −18 °C.", "C) Un dron desciende 50 metros.", "D) El automóvil está en el estacionamiento −1."], "C", "Descender describe una acción de traslado (movimiento), mientras que las otras son ubicaciones estáticas."),
                ("3. Si el punto de referencia cero representa el año de nacimiento de una civilización, ¿qué indica el número −250?", ["A) 250 años después de su nacimiento.", "B) 250 habitantes.", "C) 250 años antes de su nacimiento.", "D) Una distancia de 250 kilómetros."], "C", "Los números negativos en cronología indican tiempo transcurrido antes del punto cero (origen temporal).")
            ],
            
            # CIERRE Y METACOGNICIÓN
            "cierre": {
                "pregunta_sintesis": "En tus propias palabras, ¿por qué es tan importante definir el punto cero antes de usar números positivos y negativos?",
                "metacognicion": "¿Qué te resultó más fácil hoy: identificar la posición de un objeto o diferenciarla de un movimiento?",
                "celebracion": "¡Excelente trabajo! Has dominado la base de los números enteros: el punto de referencia."
            }
        },

        # -------------------------------------------------------------
        # CLASE 2
        # -------------------------------------------------------------
        {
            "num": 2,
            "title": "La recta numérica y orden en Z",
            "oa": "OA 01 - Clase 2 de 5",
            "duracion": "30 Minutos",
            "objetivo_adulto": "Acompañar al estudiante en la construcción de la recta numérica entera, comprendiendo el orden relativo (mayor y menor) y el criterio de posición hacia la derecha.",
            "ruta_sesion": "Desafío del cañón helado → Video Gancho → Conversación guiada → Explicación formal de la recta → Práctica de comparación → Miniquiz → Cierre metacognitivo.",
            "mentor_tip": "Ten a mano papel y lápiz por si el estudiante prefiere dibujar la recta numérica. Visualizar la línea horizontal aclara de inmediato las dudas.",
            "clima_emocional": "Valora la reflexión: 'Comparar negativos a veces confunde al principio porque parece al revés; tómate tu tiempo para mirar la recta'.",
            "situacion_intro": "Imagina una línea recta donde el 0 está al medio. A la derecha caminamos hacia los positivos (+1, +2, +3...). ¿Hacia dónde debemos caminar para encontrar los negativos?",
            "resp_esperada_sit": "Hacia la izquierda del cero.",
            "pista_socratica_sit": "Pregúntale: 'Si los positivos van a la derecha como en una regla normal, ¿cuál es la dirección opuesta al otro lado del cero?'",
            "referencia_texto": "En la recta numérica, el cero separa dos mundos: a la derecha los positivos y a la izquierda los negativos. Todo número a la derecha de otro es siempre MAYOR.",
            
            # VIDEO 1: GANCHO
            "v1_titulo": "Video 1 (Gancho H.O.O.K.): El Puente Colgante del Cañón Polar",
            "v1_dile_antes": "Fíjate en cómo están colocadas las estacas de seguridad a lo largo del cañón y qué números tienen marcados.",
            "v1_haz": "Reproducir video del Cañón Polar en la pantalla del estudiante.",
            "v1_script": [
                {"escena": "1", "tiempo": "0:00 - 0:08", "visual": "Vista aérea cinemática de un cañón nevado con una cuerda guía tensada de extremo a extremo.", "locucion": "En el corazón de un cañón polar, un equipo de rescate instala una línea de medición sobre la nieve.", "texto_pantalla": "Campamento Base Polar (0 m)"},
                {"escena": "2", "tiempo": "0:08 - 0:17", "visual": "Un explorador clava una estaca roja brillante en el centro con el rótulo '0'.", "locucion": "El poste central marca el punto cero: el refugio de abastecimiento.", "texto_pantalla": "Origen = 0"},
                {"escena": "3", "tiempo": "0:17 - 0:26", "visual": "La cámara viaja hacia el lado este (derecha): estacas verdes marcadas con +10 m, +20 m, +30 m hacia una colina soleada.", "locucion": "Hacia el este, en terreno elevado, cada marca avanza de diez en diez metros con números positivos.", "texto_pantalla": "Zona Este: +10 m, +20 m, +30 m"},
                {"escena": "4", "tiempo": "0:26 - 0:35", "visual": "La cámara gira y viaja al oeste (izquierda): estacas azules marcadas con −10 m, −20 m, −30 m descendiendo a una grieta fría.", "locucion": "Hacia el oeste, bajando hacia la grieta helada, las marcas avanzan hacia la izquierda con números negativos.", "texto_pantalla": "Zona Oeste: −10 m, −20 m, −30 m"},
                {"escena": "5", "tiempo": "0:35 - 0:43", "visual": "Dos sensores emiten luces: el sensor Alpha en −30 m y el sensor Beta en −10 m.", "locucion": "Un sensor se ubica en menos treinta metros y otro en menos diez metros.", "texto_pantalla": "Sensor A: −30 m | Sensor B: −10 m"},
                {"escena": "6", "tiempo": "0:43 - 0:51", "visual": "Plano general de la recta en la nieve: ¿Cuál sensor está más cerca del calor del campamento base?", "locucion": "¿Cuál de los dos sensores se encuentra más cerca del refugio cero y cuál está más a la derecha en la línea?", "texto_pantalla": "¿Quién está más a la derecha?"},
                {"escena": "7", "tiempo": "0:51 - 0:58", "visual": "Logotipo StudioSimple con una recta numérica animada iluminándose de izquierda a derecha.", "locucion": "En la recta numérica, estar más a la derecha cambia todo. ¡Vamos a explorarlo!", "texto_pantalla": "StudioSimple Matemática"}
            ],
            "v1_dile_despues": "Muy bien observado. Ahora vamos a ordenar estos puntos en la recta numérica.",
            
            # CONVERSACIÓN SOCRÁTICA
            "conv_items": [
                ("En la recta horizontal, ¿en qué lado del cero están los números positivos?", "A la derecha del cero.", "Recuérdale la dirección en que leemos o avanzamos normalmente."),
                ("¿Y en qué lado del cero ubicamos los números negativos?", "A la izquierda del cero.", "Pregunta: '¿Hacia qué lado retrocedemos desde el cero?'"),
                ("Entre −10 y −30, ¿cuál de los dos números está más cerca del 0?", "El −10.", "Pídele que cuente los pasos desde el cero: ¿con cuál das solo 10 pasos?"),
                ("Si un número está más a la derecha que otro en la recta, ¿es mayor o menor?", "Es mayor.", "Establece la regla de oro: hacia la derecha los números siempre CRECEN.")
            ],
            
            # VIDEO 2: FORMALIZACIÓN
            "v2_titulo": "Video 2 (Formalización): Criterio de Orden en la Recta Numérica",
            "v2_dile_antes": "Vamos a formalizar la regla matemática que nos permite comparar cualquier pareja de números enteros sin equivocarnos.",
            "v2_haz": "Reproducir video de formalización de la recta numérica.",
            "v2_script": [
                {"escena": "1", "tiempo": "0:00 - 0:10", "visual": "Una línea recta horizontal se traza en pantalla con flechas en ambos extremos y marcas uniformes.", "locucion": "La recta numérica es una línea infinita donde cada número entero tiene un lugar único y ordenado.", "texto_pantalla": "La Recta Numérica en Z"},
                {"escena": "2", "tiempo": "0:10 - 0:23", "visual": "Se ubica el 0 al centro. A la derecha aparecen en verde +1, +2, +3... A la izquierda en azul −1, −2, −3...", "locucion": "El cero es el origen. Hacia la derecha crecen los enteros positivos. Hacia la izquierda se extienden los enteros negativos.", "texto_pantalla": "← Menor (Izquierda) | Origen 0 | Mayor (Derecha) →"},
                {"escena": "3", "tiempo": "0:23 - 0:37", "visual": "Se comparan dos números negativos: −2 y −5. Una lupa resalta que −2 está más a la derecha que −5.", "locucion": "¡Esta es la regla universal! Dados dos números cualesquiera, el que esté situado más a la derecha en la recta es siempre el MAYOR.", "texto_pantalla": "−2 > −5 (porque −2 está más a la derecha)"},
                {"escena": "4", "tiempo": "0:37 - 0:49", "visual": "Ejemplos rápidos con símbolos: +4 > −100, 0 > −8, −3 < +1. Cada uno con una pequeña animación de posición.", "locucion": "Por eso, cualquier número positivo es mayor que cero, y el cero es siempre mayor que cualquier número negativo.", "texto_pantalla": "Positivo > 0 > Negativo"},
                {"escena": "5", "tiempo": "0:49 - 0:56", "visual": "Resumen gráfico con los símbolos de desigualdad: > (mayor que) y < (menor que).", "locucion": "Dominar la recta es dominar el orden de los números enteros.", "texto_pantalla": "Mayor que (>) | Menor que (<)"}
            ],
            "idea_clave": "En la recta numérica, cualquier número ubicado más a la derecha de otro es siempre MAYOR. Por ello, todo positivo es mayor que cero, y cero es mayor que cualquier negativo.",
            
            # PRÁCTICA CONJUNTA
            "practica": [
                ("Situación 1: Temperaturas de congelación", "¿Qué temperatura es más alta (más cálida): −2 °C o −7 °C? Usa el símbolo correspondiente.", "−2 °C > −7 °C (es más alta −2 °C)", "Pregunta: '¿Cuál de las dos temperaturas está más cerca de descongelarse (más cerca del cero y a la derecha)?'"),
                ("Situación 2: Comparación con el cero", "Compara 0 y −15 usando los símbolos > o <.", "0 > −15", "Recuérdale que el cero está a la derecha de todos los números negativos."),
                ("Situación 3: Ordenar un grupo de enteros", "Ordena de menor a mayor los siguientes números: +3, −8, 0, −2, +7.", "−8 < −2 < 0 < +3 < +7", "Comienza buscando el número que esté más a la izquierda de todos en la recta.")
            ],
            
            # MINIQUIZ
            "quiz": [
                ("1. ¿Cuál de las siguientes relaciones de orden es CORRECTA?", ["A) −9 > −3", "B) −5 > 0", "C) −4 > −12", "D) +2 < −8"], "C", "−4 está más a la derecha en la recta numérica que −12, por lo tanto −4 es mayor que −12."),
                ("2. Si ordenamos de MENOR a MAYOR los números −10, +5, −1, 0, ¿cuál es el orden correcto?", ["A) +5, 0, −1, −10", "B) −10, −1, 0, +5", "C) −1, −10, 0, +5", "D) 0, −1, +5, −10"], "B", "El menor es el que está más a la izquierda (−10), seguido de −1, luego 0 y finalmente +5."),
                ("3. ¿Cuál es el número entero mayor entre todos los números negativos?", ["A) −100", "B) −10", "C) −1", "D) 0"], "C", "−1 es el entero negativo que está más a la derecha en la recta (justo al lado izquierdo del 0), por lo que es el mayor de los negativos.")
            ],
            
            # CIERRE Y METACOGNICIÓN
            "cierre": {
                "pregunta_sintesis": "Si un amigo te dice que −50 es más grande que −5 porque el 50 es un número grande, ¿cómo se lo explicarías usando la recta numérica?",
                "metacognicion": "¿Te ayudó imaginar la recta horizontal para comparar los números negativos?",
                "celebracion": "¡Excelente! Ya sabes ordenar cualquier número en la recta numérica."
            }
        },

        # -------------------------------------------------------------
        # CLASE 3
        # -------------------------------------------------------------
        {
            "num": 3,
            "title": "Valor absoluto y distancias al cero",
            "oa": "OA 01 - Clase 3 de 5",
            "duracion": "30 Minutos",
            "objetivo_adulto": "Guiar al estudiante a conceptualizar el valor absoluto como una distancia geométrica pura (siempre positiva o cero) y reconocer los números opuestos o simétricos.",
            "ruta_sesion": "Desafío de los drones de rescate → Video Gancho → Conversación guiada → Explicación formal del valor absoluto → Práctica de distancias → Miniquiz → Cierre metacognitivo.",
            "mentor_tip": "Enfatiza que las distancias nunca se miden con signos negativos: un mapa nunca dice 'camine menos 5 metros'.",
            "clima_emocional": "Brinda confianza: 'El valor absoluto es como medir con una huincha: solo nos importa cuántos pasos hay, sin importar la dirección'.",
            "situacion_intro": "Si caminas 4 pasos hacia adelante desde tu silla, te moviste 4 pasos. Si caminas 4 pasos hacia atrás, ¿cuántos pasos recorriste en total de distancia?",
            "resp_esperada_sit": "4 pasos (la distancia recorrida es la misma).",
            "pista_socratica_sit": "Pregunta: '¿Tu cuerpo gastó energía para dar 4 pasos o se borraron tus pasos? La distancia siempre es una cantidad de pasos positiva'.",
            "referencia_texto": "El valor absoluto de un número entero es la distancia que hay entre ese número y el cero en la recta numérica. Se escribe entre dos barras verticales: |a|.",
            
            # VIDEO 1: GANCHO
            "v1_titulo": "Video 1 (Gancho H.O.O.K.): La Carrera de Drones Simétricos",
            "v1_dile_antes": "Observa la torre de control en el centro y fíjate en la distancia que recorre cada dron al volar en direcciones contrarias.",
            "v1_haz": "Reproducir video de los drones en la pantalla del estudiante.",
            "v1_script": [
                {"escena": "1", "tiempo": "0:00 - 0:08", "visual": "Pista de despegue futurista con una torre de control central marcada con un holograma brillante del número 0.", "locucion": "En el centro de pruebas, dos drones de inspección despegan desde la torre de control.", "texto_pantalla": "Torre de Control (Origen 0)"},
                {"escena": "2", "tiempo": "0:08 - 0:17", "visual": "El Dron Rojo vuela hacia la derecha a toda velocidad y aterriza en la plataforma +50 metros.", "locucion": "El Dron Rojo vuela hacia el este y se posiciona en la baliza más cincuenta metros.", "texto_pantalla": "Dron Rojo: Posición +50 m"},
                {"escena": "3", "tiempo": "0:17 - 0:26", "visual": "El Dron Azul vuela hacia la izquierda y aterriza exactamente en la baliza −50 metros.", "locucion": "Al mismo tiempo, el Dron Azul vuela en dirección contraria y se ubica en la baliza menos cincuenta metros.", "texto_pantalla": "Dron Azul: Posición −50 m"},
                {"escena": "4", "tiempo": "0:26 - 0:34", "visual": "Indicadores de batería y odómetro en las pantallas de telemetría de ambos drones marcando 50 m.", "locucion": "Aunque sus posiciones tienen signos opuestos, ambos motores consumieron exactamente la misma energía.", "texto_pantalla": "Energía consumida = 50 m recorridos"},
                {"escena": "5", "tiempo": "0:34 - 0:43", "visual": "Una regla holográfica mide desde el 0 al +50 (50 m) y desde el 0 al −50 (50 m). Las dos barras miden igual.", "locucion": "¿A qué distancia de la torre quedó cada dron? ¿Importa hacia qué lado volaron para saber la distancia recorrida?", "texto_pantalla": "Distancia a la torre = ¿50 m?"},
                {"escena": "6", "tiempo": "0:43 - 0:51", "visual": "Primer plano de los dos números: +50 y −50 encerrados entre dos barras mágicas | | que se transforman en 50.", "locucion": "En matemáticas, a esta distancia pura sin importar la dirección la llamamos valor absoluto.", "texto_pantalla": "|+50| = 50  y  |−50| = 50"},
                {"escena": "7", "tiempo": "0:51 - 0:58", "visual": "Logo StudioSimple con el símbolo de valor absoluto resplandeciendo en tonos turquesa.", "locucion": "¡Descubramos el poder del valor absoluto en la lección!", "texto_pantalla": "StudioSimple Matemática"}
            ],
            "v1_dile_despues": "Excelente. Ahora analicemos qué significa medir distancias al cero.",
            
            # CONVERSACIÓN SOCRÁTICA
            "conv_items": [
                ("¿A qué distancia de la torre 0 se encuentra el Dron Rojo en +50?", "A 50 metros de distancia.", "Cuenta cuántos metros hay entre 0 y 50."),
                ("¿A qué distancia de la torre 0 se encuentra el Dron Azul en −50?", "A 50 metros de distancia (también 50 m).", "Recuérdale: una distancia física siempre es una medida positiva."),
                ("¿Por qué decimos que +50 y −50 son números 'opuestos' o simétricos?", "Porque están a la misma distancia del cero pero en lados contrarios.", "Pídele que imagine doblar la recta por la mitad en el 0: ambos números coinciden."),
                ("¿Puede el valor absoluto de un número ser un valor negativo?", "No, nunca. Porque una distancia no puede ser negativa (es siempre 0 o positiva).", "Pregunta: '¿Existe una distancia de menos 3 metros en la vida real?'")
            ],
            
            # VIDEO 2: FORMALIZACIÓN
            "v2_titulo": "Video 2 (Formalización): Definición Geométrica de Valor Absoluto y Opuesto",
            "v2_dile_antes": "Veamos la definición formal de valor absoluto y cómo se escribe con símbolos matemáticos.",
            "v2_haz": "Reproducir video de formalización del valor absoluto.",
            "v2_script": [
                {"escena": "1", "tiempo": "0:00 - 0:10", "visual": "Aparece la recta numérica con el 0 al centro. Se resaltan dos flechas de igual longitud hacia ambos lados.", "locucion": "El valor absoluto de un número representa la distancia geométrica desde ese número hasta el cero.", "texto_pantalla": "Valor Absoluto = Distancia al 0"},
                {"escena": "2", "tiempo": "0:10 - 0:22", "visual": "Se escriben en pantalla las barras verticales | | envolviendo al número −7 y al número +7.", "locucion": "Se simboliza encerrando el número entre dos barras verticales. Por ejemplo, el valor absoluto de menos siete se escribe barra menos siete barra, y es igual a siete.", "texto_pantalla": "|−7| = 7  y  |+7| = 7"},
                {"escena": "3", "tiempo": "0:22 - 0:35", "visual": "Animación de una cinta métrica midiendo desde 0 a −12 (12 unidades) y desde 0 a 0 (0 unidades).", "locucion": "Como el valor absoluto mide una distancia, su resultado siempre es cero o un número positivo. ¡Nunca negativo! El valor absoluto de cero es simplemente cero.", "texto_pantalla": "|a| ≥ 0 para todo número entero"},
                {"escena": "4", "tiempo": "0:35 - 0:48", "visual": "Dos puntos simétricos en la recta: −a y +a. Se muestra cómo al sumarlos se anulan en el centro.", "locucion": "Dos números que tienen el mismo valor absoluto pero distinto signo se llaman números opuestos o simétricos. El opuesto de más cuatro es menos cuatro.", "texto_pantalla": "Opuesto de +a = −a\n|+4| = |−4| = 4"},
                {"escena": "5", "tiempo": "0:48 - 0:56", "visual": "Resumen gráfico con ejemplos destacados: |−15| = 15, |+8| = 8, opuesto de −9 = +9.", "locucion": "El valor absoluto nos permite medir magnitudes reales sin preocuparnos de la dirección.", "texto_pantalla": "Valor Absoluto y Simetría en Z"}
            ],
            "idea_clave": "El valor absoluto |a| es la distancia de un número al cero; siempre es positivo o cero. Dos números opuestos están a la misma distancia del cero pero en lados contrarios.",
            
            # PRÁCTICA CONJUNTA
            "practica": [
                ("Situación 1: Cálculo directo de distancias", "Calcula el valor absoluto de: |−18| y |+25|.", "|−18| = 18  y  |+25| = 25", "Recuerda que el valor absoluto elimina el signo y entrega la distancia positiva."),
                ("Situación 2: Identificar el opuesto aditivo", "¿Cuál es el número opuesto de −34? ¿Y cuál es su valor absoluto?", "El opuesto es +34 (o 34). El valor absoluto de ambos es 34.", "El opuesto tiene el signo contrario; el valor absoluto es la distancia sin signo."),
                ("Situación 3: Problema contextual de buceo", "Un ave vuela a +15 m sobre el mar y un pez nada a −15 m bajo el mar. ¿Cuál de los dos está más lejos de la superficie?", "Ambos están exactamente a la misma distancia: 15 metros.", "Calcula |+15| = 15 m y |−15| = 15 m.")
            ],
            
            # MINIQUIZ
            "quiz": [
                ("1. ¿Cuál es el resultado de resolver |−45|?", ["A) −45", "B) +45", "C) 0", "D) −1"], "B", "El valor absoluto mide la distancia al 0, por lo que |−45| = 45 (positivo)."),
                ("2. ¿Qué pareja de números representa números OPUESTOS en la recta numérica?", ["A) −8 y −8", "B) +10 y +20", "C) −14 y +14", "D) 0 y −5"], "C", "−14 y +14 tienen el mismo valor absoluto (14) y signos contrarios, ubicándose a la misma distancia del cero."),
                ("3. Si sabemos que |x| = 9, ¿cuáles son los posibles valores que puede tomar el número x?", ["A) Solo +9", "B) Solo −9", "C) Puede ser +9 o −9", "D) Solo 0"], "C", "Tanto |+9| como |−9| tienen una distancia de 9 unidades hasta el cero.")
            ],
            
            # CIERRE Y METACOGNICIÓN
            "cierre": {
                "pregunta_sintesis": "Si tuvieras que explicarle a alguien qué son las barras de valor absoluto | |, ¿qué ejemplo cotidiano usarías?",
                "metacognicion": "¿Te quedó clara la diferencia entre el opuesto de un número (cambiar signo) y su valor absoluto (distancia positiva)?",
                "celebracion": "¡Excelente! Has dominado el valor absoluto y la simetría en los números enteros."
            }
        },

        # -------------------------------------------------------------
        # CLASE 4
        # -------------------------------------------------------------
        {
            "num": 4,
            "title": "Adición y sustracción en Z",
            "oa": "OA 01 - Clase 4 de 5",
            "duracion": "30 Minutos",
            "objetivo_adulto": "Acompañar al estudiante a operar sumas y restas con números enteros mediante el modelo de desplazamientos en la recta y la regla de signos fundamentada.",
            "ruta_sesion": "Desafío del robot en Marte → Video Gancho → Conversación guiada → Explicación formal de suma y resta en Z → Práctica de cálculo → Miniquiz → Cierre metacognitivo.",
            "mentor_tip": "Asocia siempre sumar un positivo con 'avanzar a la derecha' y sumar un negativo con 'retroceder a la izquierda' (o agregar una deuda).",
            "clima_emocional": "Paciencia y modelamiento: 'La resta de negativos suele ser un gran salto mental; pensarla como sumar el opuesto lo hace sencillo y lógico'.",
            "situacion_intro": "Imagina que tienes $5.000 a favor (+5.000) y compras una colación que cuesta $3.000 (−3.000). ¿Cuánto dinero te queda? ¿Y si gastas $7.000?",
            "resp_esperada_sit": "Me quedan $2.000 a favor (+2.000). Si gasto $7.000, quedo debiendo $2.000 (−2.000).",
            "pista_socratica_sit": "Ayúdale a visualizar el saldo: 'Si pagas más de lo que tienes, pasas al lado de la deuda (números negativos)'.",
            "referencia_texto": "Sumar números del mismo signo junta sus valores absolutos conservando el signo. Sumar números de distinto signo resta sus valores absolutos y conserva el signo del mayor en valor absoluto.",
            
            # VIDEO 1: GANCHO
            "v1_titulo": "Video 1 (Gancho H.O.O.K.): El Rover Explorador en las Dunas de Marte",
            "v1_dile_antes": "Fíjate en las instrucciones de avance y retroceso que recibe el rover desde el centro de control terrestre.",
            "v1_haz": "Reproducir video del rover marciano en la pantalla del estudiante.",
            "v1_script": [
                {"escena": "1", "tiempo": "0:00 - 0:08", "visual": "Paisaje rojizo de Marte con montañas al fondo. Un rover robótico con paneles solares se encuentra detenido.", "locucion": "En el cráter Jezero de Marte, el rover científico inicia una secuencia programada de desplazamientos.", "texto_pantalla": "Punto de Partida: Posición 0 m"},
                {"escena": "2", "tiempo": "0:08 - 0:17", "visual": "Ruedas oruga girando hacia adelante. El rover avanza 6 metros sobre la arena roja.", "locucion": "Primera orden: Avanzar seis metros hacia el este. La telemetría marca más seis.", "texto_pantalla": "Comando 1: +6 m  →  Posición actual: +6 m"},
                {"escena": "3", "tiempo": "0:17 - 0:26", "visual": "El rover detecta una roca gigante, frena, y sus ruedas giran en reversa retrocediendo 10 metros.", "locucion": "Alerta de terreno: El rover debe retroceder diez metros para no volcar.", "texto_pantalla": "Comando 2: Retroceder 10 m (+6 + (−10))"},
                {"escena": "4", "tiempo": "0:26 - 0:34", "visual": "El rover pasa de largo por el punto 0 y queda detenido detrás del punto de partida.", "locucion": "Al retroceder diez metros habiendo avanzado seis, el rover cruza el punto cero hacia atrás.", "texto_pantalla": "¿Posición = −4 m?"},
                {"escena": "5", "tiempo": "0:34 - 0:43", "visual": "Tercera orden en la pantalla de control: Avanzar 7 metros (+7) desde su posición actual.", "locucion": "Última orden de la misión: Desde donde quedó, avanzar siete metros hacia adelante.", "texto_pantalla": "Comando 3: −4 + (+7) = ?"},
                {"escena": "6", "tiempo": "0:43 - 0:51", "visual": "El rover avanza y clava una bandera con sensores. ¿Dónde quedó exactamente respecto al origen 0?", "locucion": "¿En qué coordenada final quedó estacionado el rover tras todas sus maniobras?", "texto_pantalla": "¿Coordenada Final = +3 m?"},
                {"escena": "7", "tiempo": "0:51 - 0:58", "visual": "Logo StudioSimple animado con signos + y − combinándose armoniosamente.", "locucion": "Aprender a sumar y restar enteros es la clave para resolver cualquier misión.", "texto_pantalla": "StudioSimple Matemática"}
            ],
            "v1_dile_despues": "Excelente misión. Ahora analizaremos las reglas para sumar y restar enteros.",
            
            # CONVERSACIÓN SOCRÁTICA
            "conv_items": [
                ("Si estabas en +6 y retrocedes 10 metros (+6 + (−10)), ¿por qué terminas en una posición negativa?", "Porque retrocediste más metros de los que habías avanzado, cruzando el cero hacia atrás (−4).", "Pídele dar 6 pasos adelante y luego 10 pasos atrás: ¿dónde queda respecto al inicio?"),
                ("Desde −4, si avanzas 7 metros (−4 + 7), ¿qué operación haces entre el 7 y el 4?", "Se restan (7 − 4 = 3) y queda positivo (+3) porque avanzaste más hacia la derecha.", "Observa los signos distintos: se restan las cantidades y gana el mayor."),
                ("Si sumamos dos deudas, como (−5) + (−3), ¿el resultado es positivo o negativo?", "Es negativo (−8), porque juntas dos deudas o dos retrocesos.", "Si debes 5 y luego debes 3 más, ¿tienes dinero a favor o debes más?"),
                ("¿Cómo transformamos una resta como (+5) − (−2) en una suma fácil?", "Restar un negativo equivale a SUMAR su opuesto positivo: (+5) + (+2) = +7.", "Quitar una deuda es lo mismo que recibir dinero a favor.")
            ],
            
            # VIDEO 2: FORMALIZACIÓN
            "v2_titulo": "Video 2 (Formalización): Reglas de Adición y Sustracción en Z",
            "v2_dile_antes": "Presta atención a las dos reglas maestras de la adición y al secreto de la resta en enteros.",
            "v2_haz": "Reproducir video de formalización de operatoria en Z.",
            "v2_script": [
                {"escena": "1", "tiempo": "0:00 - 0:10", "visual": "Animación con fichas de colores: Fichas azules (+) y Fichas rojas (−).", "locucion": "Sumar y restar enteros es muy sencillo si seguimos dos principios fundamentales.", "texto_pantalla": "Adición y Sustracción en Z"},
                {"escena": "2", "tiempo": "0:10 - 0:23", "visual": "Caso 1: (+4) + (+3) = +7 y (−4) + (−3) = −7. Las fichas del mismo color se agrupan.", "locucion": "Regla 1 (Mismo signo): Si los números tienen el mismo signo, sumamos sus valores absolutos y conservamos el signo común.", "texto_pantalla": "Mismo Signo → Sumar valores y mantener signo\n(−4) + (−3) = −7"},
                {"escena": "3", "tiempo": "0:23 - 0:37", "visual": "Caso 2: (+8) + (−5) = +3 y (−9) + (+4) = −5. Fichas contrarias se cancelan en parejas.", "locucion": "Regla 2 (Distinto signo): Si tienen signos distintos, restamos el de menor valor absoluto al de mayor valor absoluto, y el resultado lleva el signo del que tiene mayor valor absoluto.", "texto_pantalla": "Distinto Signo → Restar valores y colocar signo del mayor\n(−9) + (+4) = −5"},
                {"escena": "4", "tiempo": "0:37 - 0:49", "visual": "Transformación visual de la resta: a − b = a + (−b). El signo de resta y el negativo se vuelven un gran signo +.", "locucion": "Regla de Oro de la Sustracción: Restar un número es exactamente igual que sumar su opuesto aditivo. Por ejemplo: diez menos menos tres es diez más tres, que resulta trece.", "texto_pantalla": "a − b = a + (−b)\n10 − (−3) = 10 + 3 = 13"},
                {"escena": "5", "tiempo": "0:49 - 0:56", "visual": "Resumen con cuadro de reglas y flechas direccionales en la recta numérica.", "locucion": "Con estas reglas, cualquier cálculo con números enteros se vuelve directo y exacto.", "texto_pantalla": "¡Dominio de la Operatoria en Z!"}
            ],
            "idea_clave": "Para sumar igual signo: se suman valores y se mantiene el signo. Distinto signo: se restan valores y manda el signo del mayor valor absoluto. Para restar: se suma el opuesto del sustraendo (a − b = a + (−b)).",
            
            # PRÁCTICA CONJUNTA
            "practica": [
                ("Situación 1: Suma de igual y distinto signo", "Resuelve: a) (−12) + (−8)   |   b) (−15) + (+20)", "a) −20   |   b) +5", "En a) junta deudas (12+8=20 con signo −). En b) resta 20−15=5 con signo + porque 20 es mayor."),
                ("Situación 2: Aplicación de la resta con el opuesto", "Calcula: (+7) − (+12)  y  (+4) − (−6).", "(+7) − (+12) = 7 + (−12) = −5   |   (+4) − (−6) = 4 + 6 = +10", "Convierte la resta en la suma del opuesto aditivo."),
                ("Situación 3: Variación térmica", "En la mañana la temperatura era de −3 °C y al mediodía subió 8 °C. ¿Cuál es la temperatura al mediodía?", "−3 + 8 = +5 °C", "Plantea la adición: temperatura inicial (−3) + aumento (+8).")
            ],
            
            # MINIQUIZ
            "quiz": [
                ("1. ¿Cuál es el resultado de la operación: (−18) + (+7)?", ["A) −25", "B) +11", "C) −11", "D) +25"], "C", "Al tener signos distintos, restamos 18 − 7 = 11. Como 18 tiene mayor valor absoluto y es negativo, el resultado es −11."),
                ("2. Al resolver la sustracción: (−5) − (−9), ¿qué resultado se obtiene?", ["A) −14", "B) +4", "C) −4", "D) +14"], "B", "Transformamos la resta en suma del opuesto: (−5) + (+9). Al restar 9 − 5 obtenemos +4."),
                ("3. Un buzo está a −15 metros y desciende 8 metros más. ¿Cuál es su nueva posición?", ["A) −7 m", "B) +23 m", "C) −23 m", "D) −8 m"], "C", "Descender 8 m más significa sumar otro número negativo: (−15) + (−8) = −23 metros.")
            ],
            
            # CIERRE Y METACOGNICIÓN
            "cierre": {
                "pregunta_sintesis": "¿Por qué restar un número negativo termina convirtiéndose en una suma de positivos?",
                "metacognicion": "¿Qué estrategia te sirvió más para no confundir la regla de la suma con la de la resta?",
                "celebracion": "¡Fantástico! Has conquistado las operaciones de suma y resta en los números enteros."
            }
        },

        # -------------------------------------------------------------
        # CLASE 5
        # -------------------------------------------------------------
        {
            "num": 5,
            "title": "Resolución de problemas cotidianos y síntesis",
            "oa": "OA 01 - Clase 5 de 5",
            "duracion": "30 Minutos",
            "objetivo_adulto": "Consolidar el aprendizaje del OA 01 aplicando números enteros, rectas y operaciones a problemas de la vida real (banca, temperaturas, cronología y altitudes) mediante el método de 4 pasos de Pólya.",
            "ruta_sesion": "Desafío de la Base Antártica → Video Gancho → Conversación guiada → Explicación formal de resolución de problemas → Práctica de problemas reales → Miniquiz de Maestría → Cierre metacognitivo final del OA 01.",
            "mentor_tip": "Motiva a tu hijo a leer con calma el enunciado, subrayar los datos con su signo correspondiente y escribir la respuesta completa con su unidad de medida.",
            "clima_emocional": "Celebración de logro: 'Hoy cerramos todo el tema de números enteros; mira todo lo que has avanzado desde la primera clase'.",
            "situacion_intro": "En una ciudad del sur la temperatura mínima fue de −4 °C y la máxima fue de +11 °C. ¿Cuántos grados varió (subió) la temperatura durante el día?",
            "resp_esperada_sit": "Varió 15 °C (subió 15 grados).",
            "pista_socratica_sit": "Pregunta: '¿Cuántos grados subió para llegar de −4 al 0? (4 grados). ¿Y del 0 al 11? (11 grados). En total: 4 + 11 = 15 grados'.",
            "referencia_texto": "Para resolver problemas con enteros: 1) Identificar datos y asignar signos, 2) Plantear la operación matemática, 3) Calcular con cuidado, y 4) Interpretar la respuesta en el contexto del problema.",
            
            # VIDEO 1: GANCHO
            "v1_titulo": "Video 1 (Gancho H.O.O.K.): La Misión de Suministros en la Base Antártica",
            "v1_dile_antes": "Presta atención a cómo el equipo de la Base Antártica combina registros de temperatura, combustible y presupuesto financiero.",
            "v1_haz": "Reproducir video de la Base Antártica en la pantalla del estudiante.",
            "v1_script": [
                {"escena": "1", "tiempo": "0:00 - 0:08", "visual": "Plano general de la Base Científica Antártica cubierta de nieve con turbinas eólicas y antenas parabólicas.", "locucion": "En el continente blanco, los científicos deben administrar recursos vitales bajo condiciones extremas.", "texto_pantalla": "Base Científica Antártica"},
                {"escena": "2", "tiempo": "0:08 - 0:17", "visual": "Un termómetro digital marca −18 °C al amanecer y entra un frente de viento polar.", "locucion": "A las 06:00 AM el termómetro marca menos dieciocho grados Celsius.", "texto_pantalla": "06:00 AM: −18 °C"},
                {"escena": "3", "tiempo": "0:17 - 0:26", "visual": "Al mediodía sale el sol brillante sobre el hielo y la temperatura sube doce grados.", "locucion": "Hacia el mediodía, la temperatura sube doce grados gracias a la radiación solar.", "texto_pantalla": "Mediodía: Sube +12 °C  →  ¿Temp = −6 °C?"},
                {"escena": "4", "tiempo": "0:26 - 0:34", "visual": "Pantalla bancaria de la estación: Muestra un saldo inicial de deuda de −$450.000 por repuestos.", "locucion": "En la administración, la cuenta de combustible tenía una deuda de cuatrocientos cincuenta mil pesos.", "texto_pantalla": "Cuenta Combustible: −$450.000"},
                {"escena": "5", "tiempo": "0:34 - 0:43", "visual": "Llega una transferencia de subsidio científico de +$800.000. El saldo cambia y se ilumina en verde.", "locucion": "El ministerio aprueba un aporte de ochocientos mil pesos. ¿Cuál es el nuevo saldo disponible?", "texto_pantalla": "Aporte: +$800.000  →  ¿Nuevo Saldo = +$350.000?"},
                {"escena": "6", "tiempo": "0:43 - 0:51", "visual": "Gráfica integral que reúne temperaturas, saldos y altitudes del avión de rescate.", "locucion": "Todos estos desafíos de la vida real se resuelven con exactitud usando números enteros.", "texto_pantalla": "Matemática Aplicada al Mundo Real"},
                {"escena": "7", "tiempo": "0:51 - 0:58", "visual": "Logo StudioSimple con una insignia dorada de 'Maestría en Enteros Z'.", "locucion": "¡Demuestra todo lo que has aprendido en esta gran lección final!", "texto_pantalla": "StudioSimple - Maestría OA 01"}
            ],
            "v1_dile_despues": "Excelente. Ahora resolveremos problemas combinados paso a paso.",
            
            # CONVERSACIÓN SOCRÁTICA
            "conv_items": [
                ("En la Antártica, si estábamos a −18 °C y la temperatura subió 12 °C (−18 + 12), ¿por qué sigue haciendo frío bajo cero?", "Porque el aumento (+12) no alcanzó a superar los 18 grados bajo cero, quedando en −6 °C.", "Compara los valores absolutos: 18 es mayor que 12, así que el resultado sigue siendo negativo."),
                ("Si la base debía $450.000 (−450.000) y recibe $800.000 (+800.000), ¿cómo calculamos el saldo final?", "Restamos 800.000 − 450.000 = $350.000 a favor (+350.000).", "Pagamos la deuda completa y nos sobra dinero a favor."),
                ("Para calcular la variación térmica entre una mínima de −5 °C y una máxima de +10 °C, ¿qué operación representa esa diferencia?", "Restar Máxima − Mínima: (+10) − (−5) = 10 + 5 = 15 °C de variación.", "La distancia térmica entre −5 y +10 abarca 5 grados bajo cero más 10 grados sobre cero."),
                ("¿Por qué es fundamental escribir la respuesta con palabras y unidades (ej. '$350.000 a favor' o '15 °C') en vez de solo poner un número?", "Porque le da sentido real al problema y evita confusiones sobre si es dinero a favor, deuda o temperatura.", "El número es la herramienta; la respuesta explica el significado en el mundo real.")
            ],
            
            # VIDEO 2: FORMALIZACIÓN
            "v2_titulo": "Video 2 (Formalización): El Método de 4 Pasos para Resolver Problemas en Z",
            "v2_dile_antes": "Veamos la estructura infalible de 4 pasos para resolver cualquier problema de matemáticas.",
            "v2_haz": "Reproducir video de formalización del método de resolución de problemas.",
            "v2_script": [
                {"escena": "1", "tiempo": "0:00 - 0:10", "visual": "Aparece un esquema circular con 4 cuadrantes numerados del 1 al 4 con iconos ilustrativos.", "locucion": "Para resolver con éxito cualquier problema con números enteros, aplicamos el método de cuatro pasos de George Pólya.", "texto_pantalla": "Método de 4 Pasos (Pólya)"},
                {"escena": "2", "tiempo": "0:10 - 0:22", "visual": "Paso 1: Se resalta el texto de un problema y se extraen los datos asignando signos: Ganancia (+), Deuda (−), Altura (+), Profundidad (−).", "locucion": "Paso 1: Comprender y traducir. Leemos con atención y asignamos los signos correctos a cada dato.", "texto_pantalla": "Paso 1: Identificar datos con signo (+ / −)"},
                {"escena": "3", "tiempo": "0:22 - 0:35", "visual": "Paso 2: Se escribe la expresión matemática formal en una pizarra digital (ej. Saldo = −120 + 300 − 50).", "locucion": "Paso 2: Planificar y modelar. Escribimos la operación matemática que conecta los datos con la pregunta.", "texto_pantalla": "Paso 2: Plantear la expresión matemática"},
                {"escena": "4", "tiempo": "0:35 - 0:48", "visual": "Paso 3: Cálculo paso a paso respetando las reglas de signos. Paso 4: Redacción de la respuesta contextualizada.", "locucion": "Paso 3 y 4: Ejecutar el cálculo y redactar la respuesta con su unidad de medida y sentido práctico.", "texto_pantalla": "Paso 3: Resolver | Paso 4: Interpretar respuesta"},
                {"escena": "5", "tiempo": "0:48 - 0:56", "visual": "Insignia de certificación completa del OA 01 con todos los conceptos integrados.", "locucion": "¡Felicidades! Has completado el viaje por el fascinante mundo de los números enteros.", "texto_pantalla": "¡Maestría en Números Enteros Z!"}
            ],
            "idea_clave": "Resolver problemas con enteros requiere: 1) Asignar signos a los datos, 2) Plantear la operación, 3) Calcular con rigor de signos, y 4) Redactar la respuesta contextualizada.",
            
            # PRÁCTICA CONJUNTA
            "practica": [
                ("Problema 1: Movimientos en cuenta bancaria", "Una persona tiene un saldo de $25.000. Realiza un giro en cajero de $40.000 y al día siguiente le depositan $30.000. ¿Cuál es su saldo final?", "Saldo inicial (+25.000) − 40.000 = −$15.000 (deuda). Luego: −15.000 + 30.000 = +$15.000 (saldo a favor).", "Modela paso a paso: (+25.000) + (−40.000) + (+30.000)."),
                ("Problema 2: Alturas y profundidades geográficas", "El pico de una montaña está a +2.800 m sobre el nivel del mar y una fosa marina cercana está a −1.200 m. ¿Cuál es la diferencia de altura total entre la cima y el fondo de la fosa?", "Diferencia = (+2.800) − (−1.200) = 2.800 + 1.200 = 4.000 metros de diferencia.", "Resta Cima − Fondo. Recuerda que restar un negativo suma la distancia."),
                ("Problema 3: Cronología histórica", "Un filósofo griego nació en el año 384 a.C. (−384) y vivió 62 años. ¿En qué año murió?", "Año de muerte = (−384) + (+62) = −322, es decir, en el año 322 a.C.", "Avanzar en el tiempo hacia adelante significa sumar años positivos.")
            ],
            
            # MINIQUIZ
            "quiz": [
                ("1. En un frigorífico la temperatura desciende 3 °C cada hora. Si la temperatura inicial era de 6 °C, ¿cuál será la temperatura tras 4 horas?", ["A) +12 °C", "B) −6 °C", "C) −12 °C", "D) 0 °C"], "B", "En 4 horas desciende 4 × 3 = 12 °C (−12). Entonces: (+6) + (−12) = −6 °C."),
                ("2. Un ascensor parte en el piso 8 (+8), baja 11 pisos y luego sube 5 pisos. ¿En qué piso se detiene?", ["A) Piso +2", "B) Piso −3", "C) Piso +4", "D) Piso 0"], "A", "Cálculo: (+8) − 11 = −3 (subterráneo 3). Luego: (−3) + 5 = +2 (piso 2)."),
                ("3. Un comerciante registra las siguientes ganancias y pérdidas en la semana: Lunes (+15.000), Martes (−8.000), Miércoles (−12.000), Jueves (+20.000). ¿Cuál fue su balance semanal?", ["A) Pérdida de $5.000", "B) Ganancia de $15.000", "C) Ganancia de $35.000", "D) Deuda de $20.000"], "B", "Ganancias: 15.000 + 20.000 = +35.000. Pérdidas: (−8.000) + (−12.000) = −20.000. Balance: +35.000 − 20.000 = +$15.000 (ganancia).")
            ],
            
            # CIERRE Y METACOGNICIÓN
            "cierre": {
                "pregunta_sintesis": "¿Qué fue lo más valioso que aprendiste en este módulo sobre los números enteros y cómo te servirá en tu vida diaria?",
                "metacognicion": "¿Sientes mayor seguridad para resolver operaciones y problemas con números positivos y negativos?",
                "celebracion": "¡Felicitaciones! Has completado con éxito todo el Objetivo de Aprendizaje 01 de Matemática de 7° Básico."
            }
        }
    ]

    # =========================================================================
    # RENDERIZAR CADA LECCIÓN EN EL DOCUMENTO WORD
    # =========================================================================
    for l_data in lessons_data:
        # Encabezado de la lección
        h1 = doc.add_heading(f"Lección {l_data['num']}: {l_data['title']}", level=1)
        h1.paragraph_format.space_before = Pt(18)
        h1.paragraph_format.space_after = Pt(4)
        
        p_sub = doc.add_paragraph()
        r_sub = p_sub.add_run(f"Asignatura: Matemática 7° Básico | {l_data['oa']} | Duración sugerida: {l_data['duracion']}")
        r_sub.font.size = Pt(9.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
        p_sub.paragraph_format.space_after = Pt(12)
        
        # -------------------------------------------------------------
        # PASO 1: PORTADA Y PREPARACIÓN
        # -------------------------------------------------------------
        doc.add_heading("Paso 1: Portada y Preparación del Mentor", level=2)
        add_callout_box(
            "OBJETIVO DE LA SESIÓN (PARA EL ADULTO)",
            l_data["objetivo_adulto"],
            bg_color="F0FDFA",
            border_color="0D9488",
            title_color="0F766E"
        )
        add_callout_box(
            "RUTA DE LA CLASE DE HOY",
            l_data["ruta_sesion"],
            bg_color="F8FAFC",
            border_color="64748B",
            title_color="334155"
        )
        add_callout_box(
            "RECORDATORIO DE MEDIACIÓN Y CLIMA EMOCIONAL",
            f"• Mediación: {l_data['mentor_tip']}\n• Clima Emocional: {l_data['clima_emocional']}",
            bg_color="FFFBEB",
            border_color="F59E0B",
            title_color="B45309"
        )
        
        # -------------------------------------------------------------
        # PASO 2: RUTA Y SITUACIÓN INICIAL
        # -------------------------------------------------------------
        doc.add_heading("Paso 2: Conexión y Situación Inicial", level=2)
        add_callout_box(
            "DILE AL ESTUDIANTE (SITUACIÓN INICIAL)",
            l_data["situacion_intro"],
            bg_color="EFF6FF",
            border_color="3B82F6",
            title_color="1D4ED8"
        )
        
        p_resp = doc.add_paragraph()
        p_resp.add_run("• Respuesta esperada: ").bold = True
        p_resp.add_run(l_data["resp_esperada_sit"] + "\n")
        p_resp.add_run("• Guía socrática de apoyo: ").bold = True
        p_resp.add_run(l_data["pista_socratica_sit"] + "\n")
        p_resp.paragraph_format.space_after = Pt(6)
        
        add_callout_box(
            "DILE AL ESTUDIANTE (FIJACIÓN DEL PUNTO DE REFERENCIA)",
            l_data["referencia_texto"],
            bg_color="EFF6FF",
            border_color="3B82F6",
            title_color="1D4ED8"
        )
        
        # -------------------------------------------------------------
        # PASO 3: VIDEO 1 - GANCHO MULTIMEDIA (H.O.O.K.)
        # -------------------------------------------------------------
        doc.add_heading(f"Paso 3: {l_data['v1_titulo']}", level=2)
        add_callout_box(
            "DILE ANTES DE REPRODUCIR EL VIDEO",
            l_data["v1_dile_antes"],
            bg_color="EFF6FF",
            border_color="3B82F6",
            title_color="1D4ED8"
        )
        add_callout_box(
            "HAZ (ACCIÓN EN PANTALLA)",
            l_data["v1_haz"],
            bg_color="FEF2F2",
            border_color="EF4444",
            title_color="B91C1C"
        )
        
        p_tbl_lbl1 = doc.add_paragraph()
        r_tbl1 = p_tbl_lbl1.add_run("🎬 Guión Técnico y Audiovisual del Video 1 (Gancho H.O.O.K.):")
        r_tbl1.bold = True
        r_tbl1.font.size = Pt(10)
        p_tbl_lbl1.paragraph_format.space_after = Pt(4)
        
        add_video_script_table(l_data["v1_script"])
        
        add_callout_box(
            "DILE DESPUÉS DEL VIDEO",
            l_data["v1_dile_despues"],
            bg_color="EFF6FF",
            border_color="3B82F6",
            title_color="1D4ED8"
        )
        
        # -------------------------------------------------------------
        # PASO 4: CONVERSACIÓN GUIADA (4 PREGUNTAS SOCRÁTICAS)
        # -------------------------------------------------------------
        doc.add_heading("Paso 4: Conversación Guiada (4 Preguntas Socráticas)", level=2)
        p_conv_intro = doc.add_paragraph("Realiza las siguientes 4 preguntas una a una. Escucha atentamente la respuesta y utiliza la guía socrática en caso de duda:")
        p_conv_intro.paragraph_format.space_after = Pt(6)
        
        for idx_q, (q_text, ans_text, soc_text) in enumerate(l_data["conv_items"], 1):
            p_q = doc.add_paragraph()
            r_qn = p_q.add_run(f"Pregunta {idx_q}: {q_text}\n")
            r_qn.bold = True
            r_qn.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            
            p_q.add_run("• Respuesta esperada: ").bold = True
            p_q.add_run(ans_text + "\n")
            
            r_soc = p_q.add_run("• Guía socrática: ")
            r_soc.bold = True
            r_soc.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            p_q.add_run(soc_text)
            p_q.paragraph_format.space_after = Pt(6)
            
        # -------------------------------------------------------------
        # PASO 5: VIDEO 2 - EXPLICACIÓN FORMAL
        # -------------------------------------------------------------
        doc.add_heading(f"Paso 5: {l_data['v2_titulo']}", level=2)
        add_callout_box(
            "DILE ANTES DE LA EXPLICACIÓN FORMAL",
            l_data["v2_dile_antes"],
            bg_color="EFF6FF",
            border_color="3B82F6",
            title_color="1D4ED8"
        )
        add_callout_box(
            "HAZ (ACCIÓN EN PANTALLA)",
            l_data["v2_haz"],
            bg_color="FEF2F2",
            border_color="EF4444",
            title_color="B91C1C"
        )
        
        p_tbl_lbl2 = doc.add_paragraph()
        r_tbl2 = p_tbl_lbl2.add_run("🎬 Guión Técnico y Audiovisual del Video 2 (Explicación Formal):")
        r_tbl2.bold = True
        r_tbl2.font.size = Pt(10)
        p_tbl_lbl2.paragraph_format.space_after = Pt(4)
        
        add_video_script_table(l_data["v2_script"])
        
        add_callout_box(
            "IDEA CLAVE FORMALIZADA",
            l_data["idea_clave"],
            bg_color="F0FDF4",
            border_color="22C55E",
            title_color="15803D"
        )
        
        # -------------------------------------------------------------
        # PASO 6: PRÁCTICA CONJUNTA (3 SITUACIONES)
        # -------------------------------------------------------------
        doc.add_heading("Paso 6: Práctica Conjunta (3 Situaciones Contextualizadas)", level=2)
        for idx_p, (ctx_title, prompt, exp_ans, tip) in enumerate(l_data["practica"], 1):
            p_pr = doc.add_paragraph()
            r_pr_t = p_pr.add_run(f"Situación {idx_p}: {ctx_title}\n")
            r_pr_t.bold = True
            r_pr_t.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            
            p_pr.add_run(f"• Desafío: {prompt}\n")
            p_pr.add_run(f"• Respuesta esperada: ").bold = True
            p_pr.add_run(exp_ans + "\n")
            p_pr.add_run(f"• Tip socrático de mediación: ").bold = True
            p_pr.add_run(tip)
            p_pr.paragraph_format.space_after = Pt(6)
            
        # -------------------------------------------------------------
        # PASO 7: MINIQUIZ FORMATIVO (3 PREGUNTAS CON DIAGNÓSTICO)
        # -------------------------------------------------------------
        doc.add_heading("Paso 7: Miniquiz Formativo Individual", level=2)
        p_qz_info = doc.add_paragraph("El estudiante responde en su pantalla de forma individual. Cada pregunta evalúa un criterio clave con diagnóstico de distractores:")
        p_qz_info.paragraph_format.space_after = Pt(6)
        
        for q_enunc, opts, corr_key, feedback in l_data["quiz"]:
            p_qz = doc.add_paragraph()
            p_qz.add_run(f"{q_enunc}\n").bold = True
            for opt in opts:
                p_qz.add_run(f"   {opt}\n")
            r_corr = p_qz.add_run(f"   ✓ Clave correcta: Alternativa {corr_key}\n")
            r_corr.bold = True
            r_corr.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)
            p_qz.add_run(f"   Explicación diagnóstica: {feedback}")
            p_qz.paragraph_format.space_after = Pt(6)
            
        # -------------------------------------------------------------
        # PASO 8: CIERRE ORAL Y METACOGNICIÓN
        # -------------------------------------------------------------
        doc.add_heading("Paso 8: Cierre Oral, Metacognición y Celebración", level=2)
        add_callout_box(
            "PREGUNTA DE SÍNTESIS ORAL (DILE AL ESTUDIANTE)",
            l_data["cierre"]["pregunta_sintesis"],
            bg_color="EFF6FF",
            border_color="3B82F6",
            title_color="1D4ED8"
        )
        add_callout_box(
            "PREGUNTA DE METACOGNICIÓN",
            l_data["cierre"]["metacognicion"],
            bg_color="FAF5FF",
            border_color="A855F7",
            title_color="7E22CE"
        )
        add_callout_box(
            "CELEBRACIÓN Y PUENTE A LA PRÓXIMA CLASE",
            l_data["cierre"]["celebracion"],
            bg_color="F0FDF4",
            border_color="22C55E",
            title_color="15803D"
        )
        
        doc.add_page_break()

    # Guardar documento final
    doc.save(output_path)
    print(f"Documento generado exitosamente en: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\StudioSimple - Antigravity\OA01_Matematica_7Basico_Lecciones_Completas_y_Guiones.docx"
    create_oa01_word_doc(out_file)
