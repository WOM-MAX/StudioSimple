import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    NAVY = RGBColor(28, 50, 87)       # #1C3257
    TEAL = RGBColor(18, 161, 164)     # #12A1A4
    CHARCOAL = RGBColor(40, 50, 65)   # #283241
    MUTED = RGBColor(100, 115, 135)   # #647387

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = p_title.add_run("ESTUDIOSIMPLE · ARQUITECTURA PEDAGÓGICA Y DE SOFTWARE\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(9.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = TEAL

    run_title = p_title.add_run("Estandarización y Nomenclatura Genérica del Motor de 8 Etapas")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(14)
    run_meta = p_meta.add_run("Desacoplamiento entre la capa de motor y la capa de contenidos")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(11)
    run_meta.font.italic = True
    run_meta.font.color.rgb = MUTED

    # Section 1: Resumen Ejecutivo
    p_h1 = doc.add_heading(level=1)
    p_h1.paragraph_format.space_before = Pt(12)
    p_h1.paragraph_format.space_after = Pt(6)
    r_h1 = p_h1.add_run("1. Resumen Ejecutivo")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = NAVY

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(8)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run(
        "En el prototipo final analizado (https://estudiosimple-clase1-oa1.carla-orellana.chatgpt.site), la barra de navegación del apoderado "
        "y el indicador de avance del estudiante estructuran la lección en 8 etapas. Sin embargo, las etapas 3 y 4 fueron rotuladas con nombres "
        "específicos del contenido de la Clase 1 de 7° Básico:"
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    bp1 = doc.add_paragraph(style='List Bullet')
    bp1.paragraph_format.space_after = Pt(3)
    r = bp1.add_run("Etapa 3: Rotulada como 'Recorrido' (alusión a la narrativa del submarino).")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    bp2 = doc.add_paragraph(style='List Bullet')
    bp2.paragraph_format.space_after = Pt(8)
    r = bp2.add_run("Etapa 4: Rotulada como 'Posición y movimiento' (concepto temático específico del OA1).")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    p1_concl = doc.add_paragraph()
    p1_concl.paragraph_format.space_after = Pt(12)
    r = p1_concl.add_run(
        "Este documento formaliza el diagnóstico del desacoplamiento requerido, describe la función pedagógica universal de cada etapa y "
        "propone opciones de nomenclatura genérica para asegurar que el motor sea reutilizable en todas las lecciones del currículum."
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    # Section 2: Diagnóstico
    p_h2 = doc.add_heading(level=1)
    p_h2.paragraph_format.space_before = Pt(12)
    p_h2.paragraph_format.space_after = Pt(6)
    r_h2 = p_h2.add_run("2. Diagnóstico: Acoplamiento de Contenido vs. Motor")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = NAVY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    r = p2.add_run("El diseño de una plataforma modular escalable exige una separación estricta entre dos capas:")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    # Table 1: Capas
    table1 = doc.add_table(rows=3, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    headers1 = ["Capa", "Responsabilidad Principal", "Ejemplos de Elementos"]
    widths1 = [Inches(1.5), Inches(2.6), Inches(2.4)]
    
    for i, h in enumerate(headers1):
        cell = table1.cell(0, i)
        cell.width = widths1[i]
        set_cell_background(cell, "1C3257")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    data1 = [
        ("Capa de Motor (Engine)", "Estructura universal, navegación, sincronización de estados, lógica de evaluación y telemetría.", "AdultSidebar, LessonProgress, BroadcastSync, 8 etapas inmutables."),
        ("Capa de Contenido (Data)", "Temas específicos de cada lección, guiones H.O.O.K., preguntas socráticas, videos e imágenes.", "Preguntas del submarino, datos de temperatura, opciones del miniquiz.")
    ]

    for row_idx, data in enumerate(data1, start=1):
        bg = "F5F7FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = table1.cell(row_idx, col_idx)
            cell.width = widths1[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    p_riesgos = doc.add_paragraph()
    p_riesgos.paragraph_format.space_after = Pt(4)
    r = p_riesgos.add_run("Riesgos de mantener nombres específicos en las etapas del motor:")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = NAVY

    r_items = [
        ("Falta de Escalabilidad: ", "Cada una de las 114 clases requeriría redefinir etiquetas de UI en componentes de navegación, generando acoplamiento innecesario."),
        ("Sobrecarga Cognitiva: ", "El apoderado necesita previsibilidad. Si en cada clase cambian las etapas, debe reaprender el flujo. Con nombres genéricos sabe siempre qué rol le toca ejercer."),
        ("Pérdida de Identidad Funcional: ", "La etapa 3 no existe para hablar de un 'recorrido', sino para realizar un diálogo socrático de indagación. La etapa 4 no existe para hablar de 'posición', sino para formalizar la regla conceptual.")
    ]
    for bold_txt, norm_txt in r_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        rb = bp.add_run(bold_txt)
        rb.font.name = "Arial"
        rb.font.size = Pt(9.5)
        rb.font.bold = True
        rb.font.color.rgb = CHARCOAL
        rn = bp.add_run(norm_txt)
        rn.font.name = "Arial"
        rn.font.size = Pt(9.5)
        rn.font.color.rgb = CHARCOAL

    # Section 3: Matriz de 8 Etapas
    p_h3 = doc.add_heading(level=1)
    p_h3.paragraph_format.space_before = Pt(14)
    p_h3.paragraph_format.space_after = Pt(6)
    r_h3 = p_h3.add_run("3. Función Pedagógica de las 8 Etapas Universales")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = NAVY

    table2 = doc.add_table(rows=9, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False

    headers2 = ["N°", "Etapa Prototipo V3", "Función Pedagógica Real (EstudioSimple)", "Rol del Apoderado / Sistema"]
    widths2 = [Inches(0.4), Inches(1.6), Inches(2.5), Inches(2.0)]

    for i, h in enumerate(headers2):
        cell = table2.cell(0, i)
        cell.width = widths2[i]
        set_cell_background(cell, "1C3257")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    data2 = [
        ("1", "Inicio", "Encuadre curricular, conexión inicial con la vida cotidiana y preparación del apoderado.", "Establecer objetivo y revisar ruta."),
        ("2", "Video", "Activación de curiosidad mediante narrativa audiovisual inicial (Gancho).", "Reproducir y sincronizar video inicial."),
        ("3", "Recorrido\n(Específico)", "Diálogo e Indagación: Preguntas socráticas sobre lo observado antes de dar la regla formal.", "Preguntar, escuchar y dar pistas."),
        ("4", "Posición y mov.\n(Específico)", "Formalización Teórica: Explicación estructurada (segundo video/recurso) e institucionalización.", "Guiar concepto y verificar comprensión."),
        ("5", "Práctica", "Aplicación Guiada y Razonamiento: Práctica en 3 contextos y comparación adaptativa.", "Evaluar razonamiento y desbloquear desafío."),
        ("6", "Resumen", "Estrategia y Síntesis: Modelo mental de 3 pasos y revelación progresiva de ideas clave.", "Consolidar las 3 ideas principales."),
        ("7", "Miniquiz", "Comprobación Autónoma y Refuerzo: Evaluación de 3 preguntas, revisión y recuperación.", "Monitorear y facilitar refuerzo."),
        ("8", "Cierre", "Metacognición y Transferencia: Reflexión final, proyección a próxima clase y registro.", "Cerrar sesión y registrar desempeño.")
    ]

    for row_idx, data in enumerate(data2, start=1):
        bg = "F5F7FA" if row_idx % 2 == 1 else "FFFFFF"
        if row_idx in (3, 4):
            bg = "FFF4E8"  # Highlight the problematic ones
        for col_idx, text in enumerate(data):
            cell = table2.cell(row_idx, col_idx)
            cell.width = widths2[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if row_idx in (3, 4) and col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = RGBColor(190, 80, 10)
            else:
                run.font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 4: Nomenclatura Genérica
    p_h4 = doc.add_heading(level=1)
    p_h4.paragraph_format.space_before = Pt(14)
    p_h4.paragraph_format.space_after = Pt(6)
    r_h4 = p_h4.add_run("4. Opciones de Nomenclatura Genérica")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = NAVY

    p_opA = doc.add_paragraph()
    p_opA.paragraph_format.space_after = Pt(4)
    r = p_opA.add_run("Opción A: Orientada a la Experiencia del Usuario (Recomendada)")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = TEAL

    p_opA_desc = doc.add_paragraph()
    p_opA_desc.paragraph_format.space_after = Pt(6)
    r = p_opA_desc.add_run("Utiliza términos directos orientados a la acción que realizan el apoderado y el estudiante:")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = CHARCOAL

    steps_A = [
        "1. Inicio", "2. Video inicial (o Gancho)", "3. Conversación (o Comprensión)",
        "4. Explicación (o Concepto)", "5. Práctica", "6. Resumen", "7. Miniquiz", "8. Cierre"
    ]
    p_boxA = doc.add_paragraph()
    p_boxA.paragraph_format.space_after = Pt(10)
    p_boxA.paragraph_format.left_indent = Inches(0.2)
    r = p_boxA.add_run("  ➔  ".join(steps_A))
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = NAVY

    p_opB = doc.add_paragraph()
    p_opB.paragraph_format.space_after = Pt(4)
    r = p_opB.add_run("Opción B: Terminología Pedagógica Formal")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = TEAL

    p_opB_desc = doc.add_paragraph()
    p_opB_desc.paragraph_format.space_after = Pt(6)
    r = p_opB_desc.add_run("Utiliza la taxonomía didáctica curricular formal:")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = CHARCOAL

    steps_B = [
        "1. Inicio", "2. Contextualización", "3. Indagación",
        "4. Formalización", "5. Práctica guiada", "6. Síntesis", "7. Evaluación", "8. Cierre"
    ]
    p_boxB = doc.add_paragraph()
    p_boxB.paragraph_format.space_after = Pt(12)
    p_boxB.paragraph_format.left_indent = Inches(0.2)
    r = p_boxB.add_run("  ➔  ".join(steps_B))
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = NAVY

    # Section 5: Arquitectura Técnica
    p_h5 = doc.add_heading(level=1)
    p_h5.paragraph_format.space_before = Pt(14)
    p_h5.paragraph_format.space_after = Pt(6)
    r_h5 = p_h5.add_run("5. Implementación en Código (TypeScript)")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(13)
    r_h5.font.bold = True
    r_h5.font.color.rgb = NAVY

    code_p = doc.add_paragraph()
    code_p.paragraph_format.space_after = Pt(12)
    code_p.paragraph_format.left_indent = Inches(0.2)
    code_text = (
        "// 1. Constante universal en el motor (src/types/lesson.ts):\n"
        "export const LESSON_STAGES = [\n"
        "  { step: 1, id: 'start', label: 'Inicio' },\n"
        "  { step: 2, id: 'hook_video', label: 'Video inicial' },\n"
        "  { step: 3, id: 'dialogue', label: 'Conversación' },\n"
        "  { step: 4, id: 'formalization', label: 'Explicación' },\n"
        "  { step: 5, id: 'practice', label: 'Práctica' },\n"
        "  { step: 6, id: 'summary', label: 'Resumen' },\n"
        "  { step: 7, id: 'miniquiz', label: 'Miniquiz' },\n"
        "  { step: 8, id: 'closing', label: 'Cierre' }\n"
        "] as const;\n\n"
        "// 2. Contenido inyectado desde la lección (src/data/lessons/...ts):\n"
        "// Cada lección provee sus propios subtítulos temáticos internos,\n"
        "// sin modificar los rótulos universales de las 8 etapas."
    )
    r = code_p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 45, 70)

    # Save
    output_path = r"d:\StudioSimple - Antigravity\Analisis_Estandarizacion_8_Etapas_EstudioSimple.docx"
    doc.save(output_path)
    print(f"Documento guardado exitosamente en: {output_path}")

if __name__ == "__main__":
    create_document()
