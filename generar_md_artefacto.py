# -*- coding: utf-8 -*-
import docx

doc_path = r"d:\StudioSimple - Antigravity\OA01_Matematica_7Basico_Lecciones_Completas_y_Guiones.docx"
md_path = r"C:\Users\walte\.gemini\antigravity-ide\brain\655fb28a-3386-4d26-8ad2-b20fe9d84a93\OA01_Matematica_7Basico_Lecciones_Completas_y_Guiones.md"

doc = docx.Document(doc_path)
lines = []

lines.append("# StudioSimple - Matemática 7° Básico")
lines.append("## Programa Completo de Lecciones y Guiones de Video (OA 01: Números Enteros Z)\n")

for p in doc.paragraphs:
    txt = p.text.strip()
    if not txt:
        continue
    if p.style.name.startswith('Heading 1'):
        lines.append(f"\n# {txt}\n")
    elif p.style.name.startswith('Heading 2'):
        lines.append(f"\n## {txt}\n")
    elif p.style.name.startswith('Heading 3'):
        lines.append(f"\n### {txt}\n")
    else:
        lines.append(f"{txt}\n")

# Extraer tablas
table_idx = 0
for table in doc.tables:
    table_idx += 1
    rows_data = []
    for row in table.rows:
        row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        rows_data.append(row_cells)
    
    if not rows_data:
        continue
    
    # Si es tabla de una sola celda (callout box)
    if len(rows_data[0]) == 1:
        content = rows_data[0][0]
        lines.append(f"> [!NOTE]\n> {content}\n")
    else:
        # Tabla markdown estándar
        header = rows_data[0]
        lines.append("\n| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * len(header)) + " |")
        for row in rows_data[1:]:
            lines.append("| " + " | ".join(row) + " |")
        lines.append("\n")

with open(md_path, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print("Markdown generado exitosamente en:", md_path)
