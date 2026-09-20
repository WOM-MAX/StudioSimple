import docx
import os

doc_path = r"d:\StudioSimple - Antigravity\OA01_Matematica_7Basico_Lecciones_Completas_y_Guiones.docx"
size_kb = os.path.getsize(doc_path) / 1024
doc = docx.Document(doc_path)

print(f"Archivo generado: {doc_path}")
print(f"Tamaño: {size_kb:.2f} KB")
print(f"Número de párrafos: {len(doc.paragraphs)}")
print(f"Número de tablas: {len(doc.tables)}")

headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading 1')]
print("Títulos de nivel 1 encontrados:")
for h in headings:
    print(f" - {h}")
